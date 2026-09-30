# -*- coding: utf-8 -*-
"""独立质量复核：从落盘文件重算，交叉核对ML全链路。只输出PASS/FAIL与证据。"""
import json,numpy as np,pandas as pd
from sklearn.model_selection import RepeatedKFold
D=r'F:\研究生gpt\周老师\9.9处理\2\3\4'
SRC=r'F:\研究生gpt\周老师\9.9处理\2\0909-按分数_青年脑机接口注意力训练体验与信任调查.xlsx'
ok=True
def chk(name,cond,evi=''):
    global ok;ok=ok and cond;print(('[PASS] ' if cond else '[FAIL] ')+name+('  | '+str(evi) if evi else ''))
groups={'NT':['Q10','Q11','Q15','Q16'],'ST':['Q12','Q14','Q20','Q22'],'PC':['Q13','Q19','Q23','Q27'],
'FS':['Q9','Q29','Q35','Q40'],'PFE':['Q17','Q18','Q21','Q36'],'WCE':['Q6','Q7','Q8','Q34'],
'PSC':['Q24','Q25','Q26','Q28'],'PARV':['Q30','Q31','Q32','Q33'],'RI':['Q37','Q38','Q39','Q41']}
raw=pd.read_excel(SRC);model=pd.read_csv(D+r'\建模数据集_量表分.csv')
# 1 样本量与构念均分独立重算
chk('样本量274',len(raw)==274==len(model),(len(raw),len(model)))
maxdiff=0
for g,qs in groups.items():
    mine=raw[qs].apply(pd.to_numeric,errors='coerce').mean(axis=1).values
    diff=np.nanmax(np.abs(mine-model[g].values));maxdiff=max(maxdiff,diff)
chk('9构念均分=落盘csv(逐格)',maxdiff<1e-12,'最大差%.2e'%maxdiff)
chk('36题无缺失/取值1-5',int(raw[[q for qs in groups.values() for q in qs]].isna().sum().sum())==0
    and raw[[q for qs in groups.values() for q in qs]].min().min()==1
    and raw[[q for qs in groups.values() for q in qs]].max().max()==5)
# 2 Q42对拍老师
q=pd.read_csv(D+r'\Q42_主观优先级.csv').set_index('因素')
tea={'PARV':(57,65,122),'PFE':(43,72,115),'WCE':(90,17,107),'FS':(33,52,85),'NT一致':(39,38,77),'PSC':(12,30,42)}
qok=all((int(q.loc[f,'第一选择']),int(q.loc[f,'第二选择']),int(q.loc[f,'进入前二总']))==t for f,t in tea.items())
chk('Q42六因素频次=老师表',qok,'第一/第二合计%d/%d'%(q.第一选择.sum(),q.第二选择.sum()))
# 3 外层CV无泄漏：每repeat每样本恰在test一次、50折
rkf=list(RepeatedKFold(n_splits=5,n_repeats=10,random_state=2026).split(model))
chk('外层=50折',len(rkf)==50,len(rkf))
rep_ok=True
for r in range(10):
    folds=rkf[r*5:(r+1)*5];testidx=sorted(np.concatenate([te for _,te in folds]))
    if testidx!=list(range(274)):rep_ok=False
chk('每repeat 5折恰好覆盖全部274样本(无重复/遗漏)',rep_ok)
# 4 逐折SHAP稳定性独立重算 vs json
fold=pd.read_csv(D+r'\shap_fold_stability_full.csv');R=json.load(open(D+r'\ml_nested_results_full.json',encoding='utf-8'))
chk('逐折明细=2任务×50折×6特征=600行',len(fold)==600,len(fold))
stab_ok=True
for t in['NT','ST']:
    sub=fold[fold.target==t]
    for s in R['targets'][t]['stability']:
        f=s['feat'];d=sub[sub.feat==f]
        top2=(d['rank']<=2).mean()*100
        if abs(top2-s['top2_pct'])>1.01 or abs(d.mean_abs.mean()-s['mean_abs'])>1e-3:stab_ok=False
chk('json稳定性=逐折明细独立重算',stab_ok)
# 5 json内部一致性：final排序与mean_abs降序一致；性能合理
intr=True
for t in['NT','ST']:
    fi=R['targets'][t]['final_importance'];vals=[x['mean_abs_shap'] for x in fi]
    rk=[x['fullrank'] for x in fi]
    if vals!=sorted(vals,reverse=True) or sorted(rk)!=list(range(1,7)):intr=False
    for m in('EN','XGB'):
        d=R['targets'][t]['perf'][m]
        if not(-.5<d['R2']<1 and d['RMSE']>0 and d['MAE']>0):intr=False
chk('SHAP排序连续且性能数值合理',intr)
# 6 关键结论方向自洽（PARV NT>>ST; PFE两任务都强）
nt={x['feat']:x['fullrank'] for x in R['targets']['NT']['final_importance']}
st={x['feat']:x['fullrank'] for x in R['targets']['ST']['final_importance']}
chk('PARV: NT第1且ST第5(NT特异)',nt['PARV']==1 and st['PARV']==5,(nt['PARV'],st['PARV']))
chk('PFE: 两任务都在前二(双轨共享)',nt['PFE']<=2 and st['PFE']<=2,(nt['PFE'],st['PFE']))
# 7 Excel关键值=json
import openpyxl
wb=openpyxl.load_workbook(D+r'\机器学习结果汇总_ElasticNet_XGBoost_SHAP.xlsx')
ws=wb['①模型性能CV'];xlvals={(ws.cell(r,1).value,ws.cell(r,2).value):ws.cell(r,3).value for r in range(2,6)}
xok=True
for t in['NT','ST']:
    for key,m in[('Elastic','EN'),('XGB','XGB')]:
        xl=[v for (tt,mm),v in xlvals.items() if tt==t and key in mm][0]
        if abs(xl-R['targets'][t]['perf'][m]['R2'])>1e-6:xok=False
chk('Excel性能R²=json',xok,list(xlvals))
print('\n==== 总体:', '全部通过 ✅' if ok else '存在问题 ❌====')
# 打印最终关键数字一览
print('\n性能OOF: ',{t:{m:R['targets'][t]['perf'][m]['R2'] for m in('EN','XGB')} for t in['NT','ST']})
